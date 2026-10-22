"""Shared fixtures for the hidden tests.

Everything is driven through the product's HTTP execution API (POST /v1/execute)
with the operator's sandbox directory set before the server starts. Expected
paths and clean-ups are computed here with the standard library only.
"""
import http.server
import io
import os
import shutil
import threading
import uuid
from pathlib import Path

import pytest

ROOT = Path(os.environ.get("FLYTO_GRADING_ROOT", "/tmp/flyto-grading")).resolve()
SANDBOX = ROOT / "sandbox"
OUTSIDE = ROOT / "outside"
for _d in (SANDBOX, OUTSIDE):
    _d.mkdir(parents=True, exist_ok=True)

TOKEN = "grading-token"
os.environ["FLYTO_API_TOKEN"] = TOKEN
os.environ["FLYTO_SANDBOX_DIR"] = str(SANDBOX)
os.environ.pop("FLYTO_ALLOW_ABSOLUTE_PATHS", None)
# image.download fetches from a local test server only.
os.environ["FLYTO_ALLOW_PRIVATE_NETWORK"] = "true"


def _png_bytes(color="blue", size=(8, 8)):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


class _ImageHandler(http.server.BaseHTTPRequestHandler):
    body = b""

    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-Type", "image/png")
        self.send_header("Content-Length", str(len(self.body)))
        self.end_headers()
        self.wfile.write(self.body)

    def log_message(self, *args):
        pass


_ImageHandler.body = _png_bytes()
_SERVER = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _ImageHandler)
IMAGE_PORT = _SERVER.server_address[1]
os.environ["FLYTO_HTTP_ALLOWED_PORTS"] = str(IMAGE_PORT)
threading.Thread(target=_SERVER.serve_forever, daemon=True).start()


def image_url(name="picture.png"):
    return f"http://127.0.0.1:{IMAGE_PORT}/{name}"


# --------------------------------------------------------------------------
# Browser stub for modules that need an open browser session. It answers the
# few page calls the paginator makes; the product code does every file write.
# --------------------------------------------------------------------------
class _Page:
    url = "https://example.test/list?page=1"


class StubBrowser:
    """A page with one batch of items. With fail_after_first=True the page
    disappears after the first batch (an interrupted run); otherwise the run
    finishes normally."""

    def __init__(self, fail_after_first):
        self.page = _Page()
        self.fail_after_first = fail_after_first
        self.calls = 0

    async def evaluate(self, script, arg=None):
        if "querySelectorAll(itemSelector)" in script:
            self.calls += 1
            if self.fail_after_first and self.calls > 1:
                raise RuntimeError("page went away")
            return [{"__index": 0, "text": "item", "html": "item"}]
        if self.fail_after_first:
            raise RuntimeError("page went away")
        return False

    async def goto(self, url):
        return None

    async def wait(self, *a, **k):
        return None

    async def close(self):
        pass


@pytest.fixture(scope="session")
def app_client():
    from fastapi.testclient import TestClient
    from core.api.server import create_app

    app = create_app(evidence_path=ROOT / "evidence")
    with TestClient(app) as client:
        client.flyto_app = app
        yield client


@pytest.fixture
def run(app_client):
    """Execute one module through POST /v1/execute; returns (status, json)."""

    def _run(module_id, params, context=None):
        body = {"module_id": module_id, "params": params}
        if context is not None:
            body["context"] = context
        resp = app_client.post(
            "/v1/execute", json=body, headers={"Authorization": f"Bearer {TOKEN}"}
        )
        try:
            data = resp.json()
        except ValueError:
            data = {}
        return resp.status_code, data

    return _run


@pytest.fixture
def browser_session(app_client):
    """Register a stub browser session; yields a factory taking fail_after_first."""
    sessions = app_client.flyto_app.state.server.browser_sessions
    made = []

    def _make(fail_after_first=True):
        sid = "grading-" + uuid.uuid4().hex[:8]
        sessions[sid] = StubBrowser(fail_after_first)
        made.append(sid)
        return sid

    yield _make
    for sid in made:
        sessions.pop(sid, None)


@pytest.fixture
def inside():
    """A fresh directory inside the sandbox."""
    d = SANDBOX / ("in-" + uuid.uuid4().hex[:10])
    d.mkdir(parents=True)
    yield d
    shutil.rmtree(d, ignore_errors=True)


@pytest.fixture
def outside():
    """A fresh directory outside the sandbox (a sibling of it)."""
    d = OUTSIDE / ("out-" + uuid.uuid4().hex[:10])
    d.mkdir(parents=True)
    yield d
    shutil.rmtree(d, ignore_errors=True)


def snapshot(d):
    """Every path under d, relative, with file contents."""
    out = {}
    for p in sorted(Path(d).rglob("*")):
        rel = str(p.relative_to(d))
        out[rel] = p.read_bytes() if p.is_file() and not p.is_symlink() else None
    return out


def refused(status, data):
    """The run was refused: a client error, or the API reports it as not ok."""
    if status >= 400:
        return True
    if not isinstance(data, dict) or data.get("ok") is not True:
        return True
    inner = data.get("data")
    return isinstance(inner, dict) and inner.get("ok") is False


def make_png(path, color="red", size=(40, 40)):
    from PIL import Image
    Image.new("RGB", size, color).save(str(path))
    return str(path)


def make_docx(path, with_image=None):
    from docx import Document
    d = Document()
    d.add_paragraph("hello")
    if with_image:
        d.add_picture(with_image)
    d.save(str(path))
    return str(path)


def make_pdf(path):
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path))
    c.drawString(100, 700, "hello")
    c.save()
    return str(path)
